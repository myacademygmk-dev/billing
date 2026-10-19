from __future__ import annotations

import csv
import io
import tempfile
import uuid
from datetime import date, datetime
from decimal import Decimal
from typing import Generator

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import and_, func, or_, select
from sqlalchemy.orm import Session

from app.api.deps import get_current_user
from app.core.database import get_db
from app.models.enums import StudentStatus
from app.models.payment import Payment
from app.models.student import Student
from app.models.student_fee import StudentFee
from app.models.user import User
from app.services.billing import fee_period_label, month_label, payment_period_label
from app.services.billing import normalize_month
from app.services.reporting import monthly_students_query, pending_students_query

router = APIRouter()

# Chunk size for DB-level streaming
_CHUNK_SIZE = 500
_XLSX_READ_SIZE = 1024 * 64


def _csv_row(row: list[str]) -> str:
    """Encode a single CSV row to string."""
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(row)
    return buf.getvalue()


def _csv_streaming_response(filename: str, row_generator: Generator[str, None, None]) -> StreamingResponse:
    """Stream CSV rows one at a time to the client."""
    return StreamingResponse(
        row_generator,
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/students.csv")
def export_students_csv(
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
) -> StreamingResponse:
    def generate():
        yield _csv_row([
            "student_code", "name", "class", "school", "dob", "doj", "gender",
            "contact_no", "father_name", "mother_name", "father_phone", "mother_phone",
            "whatsapp", "father_occupation", "mother_occupation",
            "hobbies", "address", "email", "fee", "status"
        ])
        last_code: str | None = None
        while True:
            stmt = (
                select(Student, func.coalesce(StudentFee.expected_fee_amount, 0).label("monthly_fee"))
                .outerjoin(StudentFee, StudentFee.student_id == Student.id)
                .order_by(Student.student_code)
                .limit(_CHUNK_SIZE)
            )
            if last_code is not None:
                stmt = stmt.where(Student.student_code > last_code)
            rows = db.execute(stmt).all()
            if not rows:
                break
            for s, monthly_fee in rows:
                last_code = s.student_code
                yield _csv_row([
                    s.student_code, s.name, s.class_name or "",
                    s.school_name or "", 
                    s.date_of_birth.isoformat() if s.date_of_birth else "",
                    s.joined_date.isoformat() if s.joined_date else "",
                    s.gender or "",
                    s.contact_no or "", s.father_name or "", s.mother_name or "",
                    s.parent_phone or "", s.parent_phone_2 or "",
                    s.whatsapp_no or "", s.father_occupation or "", s.mother_occupation or "",
                    s.hobbies or "", s.address or "", s.student_email or "",
                    str(monthly_fee), s.status.value,
                ])
            if len(rows) < _CHUNK_SIZE:
                break

    return _csv_streaming_response("students.csv", generate())


@router.get("/payments.csv")
def export_payments_csv(
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
    student_id: uuid.UUID | None = None,
    from_dt: datetime | None = Query(default=None, alias="from"),
    to_dt: datetime | None = Query(default=None, alias="to"),
) -> StreamingResponse:
    def generate():
        yield _csv_row(["receipt_no", "roll_no", "student_name", "class", "amount", "mode", "reference_no", "notes", "fee_period", "paid_at"])
        last_paid_at: datetime | None = None
        last_receipt_no: str | None = None
        while True:
            stmt = (
                select(
                    Payment,
                    Student.student_code.label("student_code"),
                    Student.name.label("student_name"),
                    Student.class_name.label("class_name"),
                )
                .join(Student, Payment.student_id == Student.id)
                .order_by(Payment.paid_at.desc(), Payment.receipt_no.desc())
                .limit(_CHUNK_SIZE)
            )
            if student_id:
                stmt = stmt.where(Payment.student_id == student_id)
            if from_dt:
                stmt = stmt.where(Payment.paid_at >= from_dt)
            if to_dt:
                stmt = stmt.where(Payment.paid_at <= to_dt)
            if last_paid_at is not None and last_receipt_no is not None:
                stmt = stmt.where(
                    or_(
                        Payment.paid_at < last_paid_at,
                        and_(Payment.paid_at == last_paid_at, Payment.receipt_no < last_receipt_no),
                    )
                )
            rows = db.execute(stmt).all()
            if not rows:
                break
            for p, student_code, student_name, class_name in rows:
                last_paid_at = p.paid_at
                last_receipt_no = p.receipt_no
                yield _csv_row([
                    p.receipt_no, student_code, student_name, class_name,
                    str(p.amount), p.mode.value if p.mode else "",
                    p.reference_no or "", p.notes or "",
                    fee_period_label(p.billing_start_month, p.billing_cycle_months) or "",
                    p.paid_at.isoformat() if p.paid_at else "",
                ])
            if len(rows) < _CHUNK_SIZE:
                break

    return _csv_streaming_response("payments.csv", generate())


@router.get("/pending.csv")
def export_pending_csv(
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
) -> StreamingResponse:
    def generate():
        yield _csv_row(["student_code", "name", "expected_fee", "paid_total", "pending"])
        base = pending_students_query().subquery()
        last_pending: Decimal | None = None
        last_code: str | None = None
        while True:
            stmt = (
                select(base)
                .order_by(base.c.pending.desc(), base.c.student_code)
                .limit(_CHUNK_SIZE)
            )
            if last_pending is not None and last_code is not None:
                stmt = stmt.where(
                    or_(
                        base.c.pending < last_pending,
                        and_(base.c.pending == last_pending, base.c.student_code > last_code),
                    )
                )
            rows = db.execute(stmt).mappings().all()
            if not rows:
                break
            for row in rows:
                last_pending = row["pending"]
                last_code = row["student_code"]
                yield _csv_row([
                    row["student_code"],
                    row["name"],
                    str(row["expected_fee"]),
                    str(row["paid_total"]),
                    str(row["pending"]),
                ])
            if len(rows) < _CHUNK_SIZE:
                break

    return _csv_streaming_response("pending.csv", generate())


@router.get("/monthly-students.csv")
def export_monthly_students_csv(
    month: date = Query(...),
    payment_state: str = Query("unpaid", pattern="^(paid|unpaid|all)$"),
    search: str | None = None,
    class_code: str | None = None,
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
) -> StreamingResponse:
    selected_month = normalize_month(month)

    def generate():
        yield _csv_row(["student_code", "name", "class_name", "payment_period", "monthly_fee", "month", "status", "receipt_no"])
        base = monthly_students_query(selected_month, payment_state, search, class_code).where(
            Student.status == StudentStatus.active
        ).subquery()
        last_code: str | None = None
        while True:
            stmt = select(base).order_by(base.c.student_code).limit(_CHUNK_SIZE)
            if last_code is not None:
                stmt = stmt.where(base.c.student_code > last_code)
            rows = db.execute(stmt).mappings().all()
            if not rows:
                break
            for row in rows:
                last_code = row["student_code"]
                is_paid = row["payment_id"] is not None
                yield _csv_row([
                    row["student_code"],
                    row["name"],
                    row["class_name"] or "",
                    payment_period_label(row["payment_period"]),
                    str(row["monthly_fee"]),
                    month_label(selected_month),
                    "paid" if is_paid else "unpaid",
                    row["receipt_no"] or "",
                ])
            if len(rows) < _CHUNK_SIZE:
                break

    return _csv_streaming_response(f"students_{payment_state}_{selected_month.isoformat()}.csv", generate())


@router.get("/students.xlsx")
def export_students_excel(
    attendance_from: date | None = Query(default=None, description="Include attendance on or after this date"),
    attendance_to: date | None = Query(default=None, description="Include attendance on or before this date"),
    marks_from: date | None = Query(default=None, description="Include marks for exams ending on or after this date"),
    marks_to: date | None = Query(default=None, description="Include marks for exams starting on or before this date"),
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
) -> StreamingResponse:
    """Export students as multi-sheet Excel: Full Details, Summary, Attendance, Marks."""
    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill, Alignment
    from app.models.attendance import StudentAttendance
    from app.models.exam import Mark, Exam

    wb = openpyxl.Workbook(write_only=True)
    header_font = Font(bold=True, color='FFFFFF', size=10)
    header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')

    def header_row(ws, values: list[str]):
        cells = []
        for value in values:
            cell = WriteOnlyCell(ws, value=value)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cells.append(cell)
        ws.append(cells)

    # --- Sheet 1: Full Details ---
    ws1 = wb.create_sheet("Full Details")
    header_row(ws1, [
        "Roll No", "Name", "Class", "School", "DOB", "DOJ", "Gender",
        "Contact", "Father Name", "Mother Name", "Father Phone", "Mother Phone",
        "WhatsApp", "Father Occupation", "Mother Occupation",
        "Hobbies", "Address", "Email", "Fee", "Status"
    ])

    student_stmt = (
        select(Student, func.coalesce(StudentFee.expected_fee_amount, 0).label("monthly_fee"))
        .outerjoin(StudentFee, StudentFee.student_id == Student.id)
        .order_by(Student.student_code)
    )
    last_code: str | None = None
    while True:
        stmt = student_stmt.limit(_CHUNK_SIZE)
        if last_code is not None:
            stmt = stmt.where(Student.student_code > last_code)
        student_rows = db.execute(stmt).all()
        if not student_rows:
            break
        for s, monthly_fee in student_rows:
            last_code = s.student_code
            ws1.append([
                s.student_code, s.name, s.class_name or "", s.school_name or "",
                s.date_of_birth.isoformat() if s.date_of_birth else "",
                s.joined_date.isoformat() if s.joined_date else "",
                s.gender or "", s.contact_no or "",
                s.father_name or "", s.mother_name or "",
                s.parent_phone or "", s.parent_phone_2 or "",
                s.whatsapp_no or "", s.father_occupation or "", s.mother_occupation or "",
                s.hobbies or "", s.address or "", s.student_email or "",
                str(monthly_fee), s.status.value,
            ])
        if len(student_rows) < _CHUNK_SIZE:
            break

    # --- Sheet 2: Summary (Short) ---
    ws2 = wb.create_sheet("Summary")
    header_row(ws2, ["Roll No", "Name", "Class", "Fee", "Status"])
    last_code = None
    while True:
        stmt = student_stmt.limit(_CHUNK_SIZE)
        if last_code is not None:
            stmt = stmt.where(Student.student_code > last_code)
        student_rows = db.execute(stmt).all()
        if not student_rows:
            break
        for s, monthly_fee in student_rows:
            last_code = s.student_code
            ws2.append([s.student_code, s.name, s.class_name or "", str(monthly_fee), s.status.value])
        if len(student_rows) < _CHUNK_SIZE:
            break

    # --- Sheet 3: Attendance ---
    ws3 = wb.create_sheet("Attendance")
    header_row(ws3, ["Roll No", "Name", "Class", "Date", "Status"])
    attendance_stmt = (
        select(StudentAttendance, Student)
        .join(Student, StudentAttendance.student_id == Student.id)
        .order_by(StudentAttendance.date.desc(), Student.student_code.desc())
    )
    if attendance_from:
        attendance_stmt = attendance_stmt.where(StudentAttendance.date >= attendance_from)
    if attendance_to:
        attendance_stmt = attendance_stmt.where(StudentAttendance.date <= attendance_to)
    last_attendance_date: date | None = None
    last_attendance_code: str | None = None
    while True:
        stmt = attendance_stmt.limit(_CHUNK_SIZE)
        if last_attendance_date is not None and last_attendance_code is not None:
            stmt = stmt.where(
                or_(
                    StudentAttendance.date < last_attendance_date,
                    and_(StudentAttendance.date == last_attendance_date, Student.student_code < last_attendance_code),
                )
            )
        attendance_rows = db.execute(stmt).all()
        if not attendance_rows:
            break
        for att, stu in attendance_rows:
            last_attendance_date = att.date
            last_attendance_code = stu.student_code
            ws3.append([stu.student_code, stu.name, stu.class_name or "", att.date.isoformat(), att.status])
        if len(attendance_rows) < _CHUNK_SIZE:
            break

    # --- Sheet 4: Marks ---
    ws4 = wb.create_sheet("Marks")
    header_row(ws4, ["Roll No", "Name", "Class", "Exam", "Marks Obtained", "Max Marks", "Grade"])
    mark_stmt = (
        select(Mark, Student, Exam)
        .join(Student, Mark.student_id == Student.id)
        .join(Exam, Mark.exam_id == Exam.id)
        .order_by(Exam.name, Student.student_code, Mark.id)
    )
    if marks_from:
        mark_stmt = mark_stmt.where(func.coalesce(Exam.end_date, Exam.start_date) >= marks_from)
    if marks_to:
        mark_stmt = mark_stmt.where(func.coalesce(Exam.start_date, Exam.end_date) <= marks_to)
    last_exam_name: str | None = None
    last_mark_student_code: str | None = None
    last_mark_id: uuid.UUID | None = None
    while True:
        stmt = mark_stmt.limit(_CHUNK_SIZE)
        if last_exam_name is not None and last_mark_student_code is not None and last_mark_id is not None:
            stmt = stmt.where(
                or_(
                    Exam.name > last_exam_name,
                    and_(Exam.name == last_exam_name, Student.student_code > last_mark_student_code),
                    and_(
                        Exam.name == last_exam_name,
                        Student.student_code == last_mark_student_code,
                        Mark.id > last_mark_id,
                    ),
                )
            )
        mark_rows = db.execute(stmt).all()
        if not mark_rows:
            break
        for m, stu, exam in mark_rows:
            last_exam_name = exam.name
            last_mark_student_code = stu.student_code
            last_mark_id = m.id
            ws4.append([
                stu.student_code, stu.name, stu.class_name or "",
                exam.name, str(m.marks_obtained), str(m.max_marks), m.grade or "",
            ])
        if len(mark_rows) < _CHUNK_SIZE:
            break

    buf = tempfile.TemporaryFile()
    wb.save(buf)
    buf.seek(0)

    def stream_file():
        try:
            while True:
                chunk = buf.read(_XLSX_READ_SIZE)
                if not chunk:
                    break
                yield chunk
        finally:
            buf.close()

    return StreamingResponse(
        stream_file(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=students_export.xlsx"},
    )
